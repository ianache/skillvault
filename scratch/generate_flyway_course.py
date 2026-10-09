import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Color Palette Config (Blue, Black, Gray, White)
COLOR_PRIMARY_RGB = (0, 86, 163)       # Blue #0056a3
COLOR_SECONDARY_RGB = (0, 0, 0)        # Black
COLOR_TERTIARY_RGB = (128, 128, 128)   # Gray
COLOR_LIGHT_GRAY_RGB = (240, 240, 240) # Light Gray for borders / tables
COLOR_WHITE_RGB = (255, 255, 255)      # White

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_run(run, font_name="Calibri", size_pt=11, bold=False, italic=False, rgb_color=COLOR_SECONDARY_RGB):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*rgb_color)

def add_header_footer(doc, code_text):
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"{code_text} | Ficha Técnica de Curso")
        format_run(hrun, size_pt=8.5, rgb_color=COLOR_TERTIARY_RGB)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run(f"Versión 1.0  |  Fecha: {datetime.now().strftime('%Y-%m-%d')}  |  Uso Interno y Académico")
        format_run(frun, size_pt=8.5, rgb_color=COLOR_TERTIARY_RGB)

def apply_table_styles(table):
    # Set thin gray borders for all internal lines of a table
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E5E5E5"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E5E5E5"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def generate_curso_ficha(output_path):
    print("Generando Ficha Técnica del Curso...")
    doc = Document()
    
    # Page setup
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    add_header_footer(doc, "MS-FLY-101")
    
    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(24)
    r_title = p_title.add_run("FICHA TÉCNICA DEL CURSO")
    format_run(r_title, size_pt=20, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    
    # Main Table
    table = doc.add_table(rows=10, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_styles(table)
    
    info_data = [
        ("CÓDIGO", "MS-FLY-101"),
        ("NOMBRE", "Control de Versiones de Base de Datos con Flyway en Arquitecturas de Microservicios con Spring Boot y Spring Cloud"),
        ("SUMILLA", "Este curso teórico-práctico de 4 horas capacita a los participantes en la integración y gestión de migraciones de esquemas de bases de datos relacionales (MySQL 5.x/8.x) utilizando Flyway en entornos distribuidos. Se abordará la configuración en proyectos Spring Boot 3.x/4.x y la coordinación de esquemas dentro de un ecosistema de microservicios con Spring Cloud, garantizando la consistencia de datos, la automatización del despliegue y las mejores prácticas para evitar conflictos en entornos de alta disponibilidad."),
        ("OBJETIVOS DE APRENDIZAJE", 
         "Al finalizar el curso, el participante será capaz de:\n"
         "1. Configurar Flyway en microservicios Spring Boot 3.x/4.x dentro de un ecosistema Spring Cloud para la gestión automática de migraciones de esquemas en MySQL 5.x/8.x.\n"
         "2. Diseñar scripts de migración versionados, repetibles y de reparación utilizando las mejores prácticas de nomenclatura de Flyway.\n"
         "3. Implementar estrategias seguras de migración ante escenarios de escalado horizontal y concurrencia (múltiples instancias de un mismo microservicio levantándose simultáneamente).\n"
         "4. Analizar y resolver conflictos de checksums u otros fallos comunes en la tabla flyway_schema_history."),
        ("HABILIDADES A DESARROLLAR", 
         "• Automatización del ciclo de vida del esquema de base de datos en arquitecturas distribuidas.\n"
         "• Gestión de consistencia de esquemas bajo el enfoque Database-per-Microservice.\n"
         "• Resolución estructurada de problemas de sincronización de esquemas bajo alta disponibilidad."),
        ("DURACIÓN", "4 horas (2 sesiones de 2 horas c/u)"),
        ("MODALIDAD", "Virtual"),
        ("PÚBLICO OBJETIVO", "Desarrolladores backend Java, ingenieros DevOps, ingenieros de software y arquitectos de soluciones que busquen implementar CI/CD de bases de datos en ecosistemas distribuidos."),
        ("PRERREQUISITOS", "Conocimiento básico/intermedio en Spring Boot, JPA/Hibernate y comandos SQL básicos en MySQL.")
    ]
    
    col_widths = [Inches(2.0), Inches(4.5)]
    
    # Let's populate the table
    # Row indices:
    # 0: CÓDIGO
    # 1: NOMBRE
    # 2: SUMILLA
    # 3: OBJETIVOS
    # 4: HABILIDADES
    # 5: DURACIÓN
    # 6: MODALIDAD
    # 7: PÚBLICO OBJETIVO
    # 8: PRERREQUISITOS
    
    for i, (label, val) in enumerate(info_data):
        row = table.rows[i]
        # Label cell
        cell_lbl = row.cells[0]
        cell_lbl.width = col_widths[0]
        set_cell_background(cell_lbl, "0056a3") # Blue header
        set_cell_margins(cell_lbl, top=100, bottom=100, left=150, right=150)
        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        format_run(r_lbl, size_pt=9.5, bold=True, rgb_color=COLOR_WHITE_RGB)
        
        # Value cell
        cell_val = row.cells[1]
        cell_val.width = col_widths[1]
        set_cell_background(cell_val, "FFFFFF" if i % 2 == 0 else "F9F9F9")
        set_cell_margins(cell_val, top=100, bottom=100, left=150, right=150)
        p_val = cell_val.paragraphs[0]
        
        # Split by newlines to respect formatting
        lines = val.split('\n')
        for j, line in enumerate(lines):
            if j > 0:
                p_val = cell_val.add_paragraph()
                p_val.paragraph_format.space_before = Pt(0)
                p_val.paragraph_format.space_after = Pt(2)
            r_val = p_val.add_run(line)
            format_run(r_val, size_pt=10, rgb_color=COLOR_SECONDARY_RGB)
            
    # Remove unused last row
    tbl = table._tbl
    tbl.remove(table.rows[9]._tr)
            
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Ficha técnica de curso guardada exitosamente en {output_path}")

if __name__ == "__main__":
    base_dir = r"C:\Users\ianache\Desktop\DATA\01-DOCUMENTOS\02-PROYECTOS\112-skillvault\MS-FLY-101"
    generate_curso_ficha(os.path.join(base_dir, "MS-FLY-101-ficha-curso.docx"))
