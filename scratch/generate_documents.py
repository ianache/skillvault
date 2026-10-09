import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pptx import Presentation
from pptx.util import Inches as PtInches, Pt as PtFont
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN

# Color Palette Config (Dark Red, Black, Gray, White)
COLOR_PRIMARY_RGB = (139, 0, 0)       # Dark Red #8B0000
COLOR_SECONDARY_RGB = (0, 0, 0)        # Black
COLOR_TERTIARY_RGB = (128, 128, 128)   # Gray
COLOR_LIGHT_GRAY_RGB = (240, 240, 240) # Light Gray for borders / tables
COLOR_WHITE_RGB = (255, 255, 255)      # White

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, color_hex="8B0000", bg_hex="F9F2F2"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=150, bottom=150, left=200, right=200)
    
    # Left border only (thick color)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor(*COLOR_SECONDARY_RGB)
    
    # Add an empty spacing paragraph after
    empty_p = doc.add_paragraph()
    empty_p.paragraph_format.space_before = Pt(0)
    empty_p.paragraph_format.space_after = Pt(6)

def format_run(run, font_name="Calibri", size_pt=11, bold=False, italic=False, rgb_color=COLOR_SECONDARY_RGB):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*rgb_color)

def add_header_footer(doc, title_text):
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"PY-101 — {title_text} | Ficha de Lección")
        format_run(hrun, size_pt=8.5, rgb_color=COLOR_TERTIARY_RGB)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Versión 1.0  |  Fecha: 2026-08-09  |  Confidencial y de Uso Interno")
        format_run(frun, size_pt=8.5, rgb_color=COLOR_TERTIARY_RGB)

def apply_table_styles(table):
    # Set thin gray borders for all internal lines of a table
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E5E5E5"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E5E5E5"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

# ==========================================
# 1. GENERATE FICHA DE LECCIÓN (.docx)
# ==========================================
def generate_ficha(output_path):
    print("Generando Ficha de Lección...")
    doc = Document()
    
    # Page setup
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    add_header_footer(doc, "Lección 1")
    
    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(24)
    r_title = p_title.add_run("FICHA TÉCNICA DE LECCIÓN")
    format_run(r_title, size_pt=20, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    
    # Info Table
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_styles(table)
    
    info_data = [
        ("CURSO", "PY-101 — Python y DuckDB para el Aseguramiento de Calidad y Consistencia de Datos"),
        ("LECCIÓN N°", "1 de 4"),
        ("CÓDIGO", "PY-101-L1"),
        ("TÍTULO", "Python, VS Code + Jupyter y pandas desde cero"),
        ("DURACIÓN", "2 horas (120 minutos)"),
        ("METODOLOGÍA", "Combinada (Demostración y Taller Práctico)")
    ]
    
    col_widths = [Inches(1.8), Inches(4.7)]
    for i, (label, val) in enumerate(info_data):
        row = table.rows[i]
        # Label cell
        cell_lbl = row.cells[0]
        cell_lbl.width = col_widths[0]
        set_cell_background(cell_lbl, "8B0000") # Primary color Dark Red background
        set_cell_margins(cell_lbl, top=100, bottom=100, left=150, right=150)
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_lbl = p_lbl.add_run(label)
        format_run(r_lbl, size_pt=9.5, bold=True, rgb_color=COLOR_WHITE_RGB)
        
        # Value cell
        cell_val = row.cells[1]
        cell_val.width = col_widths[1]
        set_cell_background(cell_val, "FFFFFF" if i % 2 == 0 else "F9F9F9")
        set_cell_margins(cell_val, top=100, bottom=100, left=150, right=150)
        p_val = cell_val.paragraphs[0]
        r_val = p_val.add_run(val)
        format_run(r_val, size_pt=10, rgb_color=COLOR_SECONDARY_RGB)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section: Objetivos
    h_obj = doc.add_paragraph()
    r_h_obj = h_obj.add_run("OBJETIVOS DE LA LECCIÓN")
    format_run(r_h_obj, size_pt=13, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h_obj.paragraph_format.space_before = Pt(12)
    h_obj.paragraph_format.space_after = Pt(6)
    
    p_intro_obj = doc.add_paragraph()
    r_intro_obj = p_intro_obj.add_run("Al finalizar esta lección, el participante será capaz de:")
    format_run(r_intro_obj, size_pt=10.5, italic=True)
    p_intro_obj.paragraph_format.space_after = Pt(6)
    
    objectives = [
        "Configurar VS Code con las extensiones de Python y Jupyter para iniciar un entorno interactivo local de desarrollo de manera exitosa.",
        "Ejecutar bloques de código en Jupyter Notebooks utilizando la sintaxis básica de Python (variables, tipos de datos elementales y estructuras de control).",
        "Cargar archivos estructurados (CSV, Excel, JSON) usando la librería pandas y describir la forma, dimensiones y tipos de datos del DataFrame resultante."
    ]
    for obj in objectives:
        p_bullet = doc.add_paragraph(style='List Bullet')
        p_bullet.paragraph_format.space_after = Pt(4)
        p_bullet.paragraph_format.left_indent = Inches(0.4)
        run_b = p_bullet.add_run(obj)
        format_run(run_b, size_pt=10.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section: Habilidades
    h_hab = doc.add_paragraph()
    r_h_hab = h_hab.add_run("HABILIDADES A DESARROLLAR")
    format_run(r_h_hab, size_pt=13, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h_hab.paragraph_format.space_before = Pt(12)
    h_hab.paragraph_format.space_after = Pt(6)
    
    skills = [
        "Configuración e inicialización ágil de un entorno interactivo VS Code + Jupyter.",
        "Declaración y manipulación de variables y tipos de datos básicos en Python en un entorno de celdas.",
        "Importación eficiente de datasets multi-formato a través de comandos nativos de pandas.",
        "Inspección visual y diagnósticos estructurales rápidos de DataFrames de pandas mediante propiedades base."
    ]
    for sk in skills:
        p_bullet = doc.add_paragraph(style='List Bullet')
        p_bullet.paragraph_format.space_after = Pt(4)
        p_bullet.paragraph_format.left_indent = Inches(0.4)
        run_b = p_bullet.add_run(sk)
        format_run(run_b, size_pt=10.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section: Agenda
    h_agenda = doc.add_paragraph()
    r_h_agenda = h_agenda.add_run("AGENDA DE LA SESIÓN")
    format_run(r_h_agenda, size_pt=13, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h_agenda.paragraph_format.space_before = Pt(12)
    h_agenda.paragraph_format.space_after = Pt(6)
    
    agenda_table = doc.add_table(rows=6, cols=3)
    agenda_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_styles(agenda_table)
    
    agenda_headers = ["ID", "Tema o Actividad de Aprendizaje", "Duración"]
    agenda_rows = [
        ("1", "Introducción al curso y configuración de entorno (VS Code + Jupyter + Kernel Python)", "30 min"),
        ("2", "Introducción a Jupyter Notebooks y sintaxis elemental de Python", "30 min"),
        ("3", "Conceptos clave de pandas: Series, DataFrames e Indexación básica", "20 min"),
        ("4", "Taller práctico: Carga e inspección básica de archivos CSV, Excel y JSON", "30 min"),
        ("5", "Resumen de aprendizajes, discusión de mejores prácticas y cierre de sesión", "10 min")
    ]
    
    col_widths_agenda = [Inches(0.6), Inches(4.7), Inches(1.2)]
    
    # Header row
    hdr_row = agenda_table.rows[0]
    for j, name in enumerate(agenda_headers):
        cell = hdr_row.cells[j]
        cell.width = col_widths_agenda[j]
        set_cell_background(cell, "8B0000")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(name)
        format_run(run, size_pt=10, bold=True, rgb_color=COLOR_WHITE_RGB)
        
    # Content rows
    for i, data in enumerate(agenda_rows):
        row = agenda_table.rows[i+1]
        for j, val in enumerate(data):
            cell = row.cells[j]
            cell.width = col_widths_agenda[j]
            set_cell_background(cell, "FFFFFF" if i % 2 == 0 else "F9F9F9")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            format_run(run, size_pt=9.5, rgb_color=COLOR_SECONDARY_RGB)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section: Recursos
    h_rec = doc.add_paragraph()
    r_h_rec = h_rec.add_run("RECURSOS REQUERIDOS")
    format_run(r_h_rec, size_pt=13, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h_rec.paragraph_format.space_before = Pt(12)
    h_rec.paragraph_format.space_after = Pt(6)
    
    recursos = [
        "Editor de código: Visual Studio Code instalado localmente.",
        "Entorno de Python: Versión 3.10 o superior (instalada en el path del sistema).",
        "Extensiones de VS Code: Extensión oficial de Python (Microsoft) y extensión de Jupyter (Microsoft).",
        "Librerías Python: pandas 2.0+ instalada vía gestor de paquetes (`pip`).",
        "Conjuntos de datos para taller: Un archivo CSV de clientes, un Excel de órdenes de compra y un JSON de logs de transacciones."
    ]
    for rec in recursos:
        p_bullet = doc.add_paragraph(style='List Bullet')
        p_bullet.paragraph_format.space_after = Pt(4)
        p_bullet.paragraph_format.left_indent = Inches(0.4)
        run_b = p_bullet.add_run(rec)
        format_run(run_b, size_pt=10.5)
        
    # Save document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Ficha de lección guardada exitosamente en {output_path}")

# ==========================================
# 2. GENERATE PRESENTACIÓN PPT (.pptx)
# ==========================================
def generate_presentation(output_path):
    print("Generando Presentación PPT...")
    prs = Presentation()
    prs.slide_width = PtInches(13.333) # 16:9 Aspect Ratio
    prs.slide_height = PtInches(7.5)
    
    blank_layout = prs.slide_layouts[6] # Blank slide layout
    
    slides_data = [
        # Slide 1: Portada
        {
            "type": "portada",
            "title": "Python, VS Code + Jupyter y pandas desde cero",
            "subtitle": "CURSO: PY-101 — Python y DuckDB para el Aseguramiento de Calidad y Consistencia de Datos",
            "footer": "Lección 1 de 4  |  Código: PY-101-L1",
            "notes": "Bienvenidos a la Lección 1 del curso. En esta sesión introductoria configuraremos nuestro entorno de trabajo usando VS Code y Jupyter, y cargaremos nuestros primeros archivos estructurados usando pandas. Este es el primer paso para dominar el aseguramiento de calidad de datos mediante código."
        },
        # Slide 2: Agenda
        {
            "type": "content",
            "title": "Agenda de la Sesión",
            "bullets": [
                "01. Introducción al curso y configuración de entorno (VS Code + Jupyter + Python) [30 min]",
                "02. Fundamentos de Jupyter Notebooks y sintaxis básica de Python [30 min]",
                "03. Conceptos centrales de pandas: Series y DataFrames [20 min]",
                "04. Taller práctico: Carga e inspección básica de archivos CSV, Excel y JSON [30 min]",
                "05. Resumen, mejores prácticas y cierre de sesión [10 min]"
            ],
            "notes": "Nuestra agenda de hoy está dividida de forma balanceada. Dedicaremos la primera hora a asegurar que todos tengan su entorno listo y comprendan los principios interactivos del notebook. La segunda hora estará enfocada enteramente en la librería pandas, culminando con un taller de carga práctica."
        },
        # Slide 3: Objetivos
        {
            "type": "content",
            "title": "Objetivos de Aprendizaje",
            "bullets": [
                "Configurar un entorno local de trabajo en VS Code utilizando extensiones interactivas.",
                "Utilizar celdas de Jupyter Notebook para escribir y ejecutar código Python simple y dinámico.",
                "Escribir comandos de pandas para importar datos estructurados desde tres formatos comunes.",
                "Describir la estructura fundamental de un DataFrame evaluando filas, columnas y tipos de datos."
            ],
            "notes": "Al finalizar las dos horas, deberán tener la capacidad de ejecutar un notebook, redactar bloques interactivos básicos, e importar datasets en formatos comunes. Estos son los cimientos técnicos de toda auditoría de datos posterior."
        },
        # Slide 4: Configuración del Entorno
        {
            "type": "content",
            "title": "Configuración del Entorno interactivo",
            "bullets": [
                "Visual Studio Code: El editor de código principal donde organizaremos nuestros proyectos.",
                "Extensión Python & Jupyter: Proporcionan soporte nativo para ejecutar celdas interactivas en VS Code.",
                "Python Kernel: El motor de ejecución local de Python que compilará nuestro código en tiempo real.",
                "Ventajas de Jupyter para QA/BI: Permite una exploración iterativa y visual, guardando estado de las variables y facilitando la documentación inmediata."
            ],
            "notes": "La combinación de VS Code y Jupyter es ideal para analistas de QA y BI. A diferencia de un script de Python tradicional que corre de principio a fin, Jupyter guarda las variables en memoria, lo cual permite realizar consultas de datos interactivas y rápidas paso a paso."
        },
        # Slide 5: Sintaxis Básica de Python
        {
            "type": "content",
            "title": "Fundamentos de Python en Jupyter",
            "bullets": [
                "Celdas de Código vs. Celdas Markdown: Las primeras ejecutan lógica, las segundas documentan.",
                "Tipos de Datos Elementales: Cadenas (str), Números (int, float) y Booleanos (bool).",
                "Estructura de Datos Básica: Listas y Diccionarios como contenedores clave.",
                "Sintaxis limpia: No requiere puntos y comas; la indentación define los bloques de código.",
                "Ejecución Interactiva: Shift + Enter ejecuta la celda actual y pasa a la siguiente."
            ],
            "notes": "Escribir código en Python es sumamente legible. En los notebooks utilizaremos Markdown para rotular y documentar nuestras pruebas de calidad, y celdas de código para realizar las comprobaciones lógicas del dataset."
        },
        # Slide 6: pandas: Series y DataFrames
        {
            "type": "content",
            "title": "Conceptos Clave de pandas",
            "bullets": [
                "pandas: La librería de Python estándar para manipulación y análisis de datos estructurados.",
                "Series: Estructura unidimensional indexada (similar a una columna de Excel).",
                "DataFrames: Estructura bidimensional con filas y columnas indexadas (equivalente a una tabla).",
                "Inspección Inicial: Propiedades fundamentales como DataFrame.shape y DataFrame.dtypes.",
                "Importancia en Calidad: Convierte datos brutos en objetos estructurados sobre los que correr validaciones."
            ],
            "notes": "pandas es el motor central del curso. Todo archivo que importemos se convertirá en un DataFrame, el cual tiene métodos optimizados para buscar duplicados, filtrar nulos o castear tipos de datos de forma sumamente veloz."
        },
        # Slide 7: Taller Práctico: Carga de Datos
        {
            "type": "content",
            "title": "Taller Práctico: Carga e Inspección",
            "bullets": [
                "Paso 1: Inicializar el notebook e importar la librería con 'import pandas as pd'.",
                "Paso 2: Cargar un archivo CSV mediante el método 'pd.read_csv('ruta_archivo.csv')'.",
                "Paso 3: Cargar un archivo Excel utilizando 'pd.read_excel('ruta_archivo.xlsx')'.",
                "Paso 4: Importar logs transaccionales JSON vía 'pd.read_json('ruta_archivo.json')'.",
                "Paso 5: Ejecutar DataFrame.head() y DataFrame.info() para inspeccionar la importación."
            ],
            "notes": "En la guía de laboratorio, pasaremos a la práctica. Realizaremos la importación de tres fuentes de datos diferentes e implementaremos la primera inspección visual de las tablas para validar si la carga ocurrió correctamente."
        },
        # Slide 8: Resumen y Takeaways
        {
            "type": "content",
            "title": "Puntos Clave (Takeaways)",
            "bullets": [
                "Entorno Configurado: Una estación de trabajo bien instalada es fundamental para trabajar sin fricciones.",
                "Interactividad: Los Jupyter Notebooks permiten auditar datos de forma iterativa y documentada.",
                "pandas como estándar: La carga estructurada con read_csv/excel/json nos da acceso al análisis de datos.",
                "Inspección Básica: head() e info() son nuestros primeros filtros de diagnóstico rápido de calidad."
            ],
            "notes": "Como resumen, recuerden que la clave de esta sesión es la comodidad en el entorno. Una vez dominada la carga y la inspección inicial con pandas, el diagnóstico y la transformación de datos de las siguientes sesiones se facilitarán de gran manera."
        },
        # Slide 9: Referencias y Recursos
        {
            "type": "content",
            "title": "Referencias y Recursos de Aprendizaje",
            "bullets": [
                "Documentación Oficial de pandas: pandas.pydata.org",
                "Guía de Configuración de VS Code Jupyter: code.visualstudio.com/docs/datascience/jupyter-notebooks",
                "Repositorio Oficial del Curso: [URL-PENDIENTE: repositorio del proyecto — proveer al instructor]",
                "Python para Análisis de Datos (Wes McKinney): Libro recomendado de consulta rápida."
            ],
            "notes": "Aquí disponen de enlaces clave. He colocado un placeholder para el repositorio del curso donde el instructor les proveerá los datasets y notebooks estructurados para las siguientes sesiones de práctica."
        }
    ]
    
    for slide_info in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # Slide Speaker Notes
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = slide_info.get("notes", "")
        
        if slide_info["type"] == "portada":
            # --- PORTADA DESIGN ---
            # Header dark red accent band
            shape_band = slide.shapes.add_shape(
                1, # rectangle
                0, 0, prs.slide_width, PtInches(1.8)
            )
            shape_band.fill.solid()
            shape_band.fill.fore_color.rgb = PptRGBColor(*COLOR_PRIMARY_RGB)
            shape_band.line.fill.background()
            
            # Title text box inside or just below the band
            txBox = slide.shapes.add_textbox(PtInches(0.8), PtInches(2.2), PtInches(11.7), PtInches(3.5))
            tf = txBox.text_frame
            tf.word_wrap = True
            
            # Title
            p_title = tf.paragraphs[0]
            p_title.text = slide_info["title"]
            p_title.font.name = "Calibri"
            p_title.font.size = PtFont(40)
            p_title.font.bold = True
            p_title.font.color.rgb = PptRGBColor(*COLOR_PRIMARY_RGB)
            p_title.space_after = PtFont(16)
            
            # Subtitle
            p_sub = tf.add_paragraph()
            p_sub.text = slide_info["subtitle"]
            p_sub.font.name = "Calibri"
            p_sub.font.size = PtFont(16)
            p_sub.font.bold = True
            p_sub.font.color.rgb = PptRGBColor(*COLOR_SECONDARY_RGB)
            p_sub.space_after = PtFont(24)
            
            # Footer / Code
            p_foot = tf.add_paragraph()
            p_foot.text = slide_info["footer"]
            p_foot.font.name = "Calibri"
            p_foot.font.size = PtFont(13)
            p_foot.font.italic = True
            p_foot.font.color.rgb = PptRGBColor(*COLOR_TERTIARY_RGB)
            
        else:
            # --- CONTENT SLIDE DESIGN ---
            # Red left vertical accent bar
            shape_bar = slide.shapes.add_shape(
                1, # rectangle
                0, 0, PtInches(0.4), prs.slide_height
            )
            shape_bar.fill.solid()
            shape_bar.fill.fore_color.rgb = PptRGBColor(*COLOR_PRIMARY_RGB)
            shape_bar.line.fill.background()
            
            # Title Text Box
            title_box = slide.shapes.add_textbox(PtInches(0.8), PtInches(0.5), PtInches(11.7), PtInches(1.2))
            tf_title = title_box.text_frame
            tf_title.word_wrap = True
            p_t = tf_title.paragraphs[0]
            p_t.text = slide_info["title"]
            p_t.font.name = "Calibri"
            p_t.font.size = PtFont(32)
            p_t.font.bold = True
            p_t.font.color.rgb = PptRGBColor(*COLOR_PRIMARY_RGB)
            
            # Bullets Text Box
            content_box = slide.shapes.add_textbox(PtInches(1.0), PtInches(1.8), PtInches(11.3), PtInches(5.0))
            tf_content = content_box.text_frame
            tf_content.word_wrap = True
            
            for index, bullet in enumerate(slide_info["bullets"]):
                if index == 0:
                    p_b = tf_content.paragraphs[0]
                else:
                    p_b = tf_content.add_paragraph()
                    
                p_b.text = "• " + bullet
                p_b.font.name = "Calibri"
                p_b.font.size = PtFont(15)
                p_b.font.color.rgb = PptRGBColor(*COLOR_SECONDARY_RGB)
                p_b.space_after = PtFont(12)
                p_b.level = 0
                
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentación PPT guardada exitosamente en {output_path}")

# ==========================================
# 3. GENERATE GUÍA DE LABORATORIO (.docx)
# ==========================================
def generate_lab_guide(output_path):
    print("Generando Guía de Laboratorio...")
    doc = Document()
    
    # Page setup
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    add_header_footer(doc, "Guía de Lab 1")
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(24)
    r_title = p_title.add_run("GUÍA DE LABORATORIO PRÁCTICO")
    format_run(r_title, size_pt=18, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    
    # Lab Info Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_styles(table)
    
    info_data = [
        ("CÓDIGO DE LECCIÓN", "PY-101-L1"),
        ("TÍTULO LECCIÓN", "Python, VS Code + Jupyter y pandas desde cero"),
        ("TÍTULO LABORATORIO", "Lab 1: Configuración de Entorno e Importación de Datasets Multi-formato"),
        ("DURACIÓN ESTIMADA", "30 minutos"),
        ("VERSIÓN / FECHA", "1.0  /  2026-08-09")
    ]
    
    col_widths = [Inches(2.2), Inches(4.3)]
    for i, (label, val) in enumerate(info_data):
        row = table.rows[i]
        cell_lbl = row.cells[0]
        cell_lbl.width = col_widths[0]
        set_cell_background(cell_lbl, "000000") # Dark secondary background for laboratory
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        format_run(r_lbl, size_pt=9, bold=True, rgb_color=COLOR_WHITE_RGB)
        
        cell_val = row.cells[1]
        cell_val.width = col_widths[1]
        set_cell_background(cell_val, "FFFFFF" if i % 2 == 0 else "F9F9F9")
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
        p_val = cell_val.paragraphs[0]
        r_val = p_val.add_run(val)
        format_run(r_val, size_pt=9.5, rgb_color=COLOR_SECONDARY_RGB)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # 1. OBJETIVOS DEL LABORATORIO
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. OBJETIVOS DEL LABORATORIO")
    format_run(r_h1, size_pt=12, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    
    objectives = [
        "Verificar la correcta instalación local de Python, VS Code y la extensión Jupyter.",
        "Inicializar de forma práctica un notebook interactivo (.ipynb) y ejecutar celdas de código de prueba.",
        "Importar un conjunto de datos CSV de clientes y auditar su forma y campos principales.",
        "Importar transacciones financieras desde JSON y órdenes de compra en formato Excel (.xlsx) de manera exitosa."
    ]
    for obj in objectives:
        p_bullet = doc.add_paragraph(style='List Bullet')
        p_bullet.paragraph_format.space_after = Pt(4)
        p_bullet.paragraph_format.left_indent = Inches(0.4)
        run_b = p_bullet.add_run(obj)
        format_run(run_b, size_pt=10)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # 2. ENTORNO Y PRERREQUISITOS
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. ENTORNO Y REQUISITOS TÉCNICOS")
    format_run(r_h2, size_pt=12, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    
    p_env = doc.add_paragraph()
    r_env = p_env.add_run("Para completar este laboratorio de forma exitosa, el estudiante requiere disponer de:")
    format_run(r_env, size_pt=10, italic=True)
    p_env.paragraph_format.space_after = Pt(6)
    
    env_items = [
        "Estación de trabajo con sistema operativo Windows, macOS o Linux.",
        "Python 3.10 o superior debidamente instalado y agregado al PATH de comandos del sistema.",
        "Visual Studio Code con la extensión de Python y Jupyter oficialmente instaladas.",
        "Conectividad a Internet (únicamente para la descarga inicial de librerías mediante pip)."
    ]
    for item in env_items:
        p_bullet = doc.add_paragraph(style='List Bullet')
        p_bullet.paragraph_format.space_after = Pt(4)
        p_bullet.paragraph_format.left_indent = Inches(0.4)
        run_b = p_bullet.add_run(item)
        format_run(run_b, size_pt=10)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # 3. PREPARACIÓN DEL ENTORNO
    h3 = doc.add_paragraph()
    r_h3 = h3.add_run("3. PREPARACIÓN DEL ENTORNO")
    format_run(r_h3, size_pt=12, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    
    steps_prep = [
        "Abrir una terminal o consola de comandos en su equipo.",
        "Verificar que Python esté correctamente instalado corriendo el siguiente comando:",
        "    python --version",
        "Instalar la librería de manipulación de datos pandas corriendo en su consola:",
        "    pip install pandas openpyxl",
        "Crear una carpeta de proyecto llamada 'PY-101-LAB1' en su equipo local y abrirla dentro de Visual Studio Code."
    ]
    for step in steps_prep:
        if step.strip().startswith("python") or step.strip().startswith("pip"):
            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Inches(0.6)
            p_code.paragraph_format.space_after = Pt(4)
            run_c = p_code.add_run(step)
            format_run(run_c, font_name="Consolas", size_pt=9.5, rgb_color=COLOR_PRIMARY_RGB)
        else:
            p_step = doc.add_paragraph(style='List Number')
            p_step.paragraph_format.space_after = Pt(4)
            p_step.paragraph_format.left_indent = Inches(0.4)
            run_s = p_step.add_run(step)
            format_run(run_s, size_pt=10)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # 4. EJERCICIOS PRÁCTICOS
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("4. EJERCICIOS Y PASOS DE EJECUCIÓN")
    format_run(r_h4, size_pt=12, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(6)
    
    # Ejercicio 1
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    h_ex1 = doc.add_paragraph()
    r_ex1 = h_ex1.add_run("Ejercicio 1: Creación del Jupyter Notebook e Inicialización de pandas")
    format_run(r_ex1, size_pt=11, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h_ex1.paragraph_format.space_after = Pt(4)
    
    p_ex1_desc = doc.add_paragraph()
    p_ex1_desc.paragraph_format.space_after = Pt(6)
    r_ex1_desc = p_ex1_desc.add_run("El objetivo de este ejercicio es inicializar un Jupyter Notebook en VS Code, asegurarse de que el Kernel de Python se vincule correctamente e importar pandas.")
    format_run(r_ex1_desc, size_pt=10)
    
    steps_ex1 = [
        "Dentro de VS Code, presione 'Ctrl + Shift + P' para abrir la paleta de comandos y escriba 'Create: New Jupyter Notebook', o cree un nuevo archivo llamado 'mi_primer_auditoria.ipynb'.",
        "En la parte superior derecha de la interfaz del notebook en VS Code, de clic en 'Select Kernel' y elija la instalación de Python local que preparó en el paso 3.",
        "Cree una celda de código, escriba la instrucción de importación de la librería y ejecútela presionando Shift + Enter:",
        "    import pandas as pd",
        "    print('pandas version:', pd.__version__)"
    ]
    for step in steps_ex1:
        if step.strip().startswith("import") or step.strip().startswith("print"):
            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Inches(0.6)
            p_code.paragraph_format.space_after = Pt(4)
            run_c = p_code.add_run(step)
            format_run(run_c, font_name="Consolas", size_pt=9.5, rgb_color=COLOR_PRIMARY_RGB)
        else:
            p_step = doc.add_paragraph(style='List Number')
            p_step.paragraph_format.space_after = Pt(4)
            p_step.paragraph_format.left_indent = Inches(0.4)
            run_s = p_step.add_run(step)
            format_run(run_s, size_pt=10)
            
    add_callout(doc, "✓ CHECKPOINT: Al ejecutar la celda, el notebook no debe reportar ningún error de importación y debe imprimir la versión de pandas instalada (ej. 'pandas version: 2.x.x').")
    
    # Ejercicio 2
    h_ex2 = doc.add_paragraph()
    r_ex2 = h_ex2.add_run("Ejercicio 2: Importación de un Dataset CSV de Clientes")
    format_run(r_ex2, size_pt=11, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h_ex2.paragraph_format.space_before = Pt(12)
    h_ex2.paragraph_format.space_after = Pt(4)
    
    p_ex2_desc = doc.add_paragraph()
    p_ex2_desc.paragraph_format.space_after = Pt(6)
    r_ex2_desc = p_ex2_desc.add_run("Procederemos a simular la importación de un archivo CSV local conteniendo registros de clientes para verificar su consistencia de campos inicial.")
    format_run(r_ex2_desc, size_pt=10)
    
    steps_ex2 = [
        "Asegúrese de guardar en la carpeta de su proyecto el archivo de muestra 'clientes.csv' proveído por el instructor.",
        "Cree una nueva celda en su notebook y escriba la sintaxis de carga:",
        "    # Cargar archivo CSV",
        "    df_clientes = pd.read_csv('clientes.csv')",
        "    # Mostrar las primeras 5 filas",
        "    df_clientes.head()",
        "Escriba otra celda para auditar el volumen total de filas e inspeccionar de forma preliminar las columnas y tipos de datos:",
        "    print('Dimensiones del dataset:', df_clientes.shape)",
        "    df_clientes.info()"
    ]
    for step in steps_ex2:
        if step.strip().startswith("#") or step.strip().startswith("df_") or step.strip().startswith("print") or step.strip().startswith("import"):
            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Inches(0.6)
            p_code.paragraph_format.space_after = Pt(4)
            run_c = p_code.add_run(step)
            format_run(run_c, font_name="Consolas", size_pt=9.5, rgb_color=COLOR_PRIMARY_RGB)
        else:
            p_step = doc.add_paragraph(style='List Number')
            p_step.paragraph_format.space_after = Pt(4)
            p_step.paragraph_format.left_indent = Inches(0.4)
            run_s = p_step.add_run(step)
            format_run(run_s, size_pt=10)
            
    add_callout(doc, "✓ CHECKPOINT: La ejecución del método .info() desplegará en pantalla el listado de columnas del dataset, el conteo de valores no nulos y la tipificación lógica de pandas (ej. int64, object).")
    
    # Ejercicio 3
    h_ex3 = doc.add_paragraph()
    r_ex3 = h_ex3.add_run("Ejercicio 3: Carga Avanzada de Excel y JSON")
    format_run(r_ex3, size_pt=11, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h_ex3.paragraph_format.space_before = Pt(12)
    h_ex3.paragraph_format.space_after = Pt(4)
    
    p_ex3_desc = doc.add_paragraph()
    p_ex3_desc.paragraph_format.space_after = Pt(6)
    r_ex3_desc = p_ex3_desc.add_run("Los analistas de QA deben verificar datos que a menudo vienen en múltiples formatos. En esta etapa cargará registros de Excel (.xlsx) y logs en formato JSON.")
    format_run(r_ex3_desc, size_pt=10)
    
    steps_ex3 = [
        "Guarde los archivos 'ordenes_compra.xlsx' y 'logs_transacciones.json' en el directorio de trabajo.",
        "Cargue las órdenes de compra utilizando el motor openpyxl en una celda de código dedicada:",
        "    df_ordenes = pd.read_excel('ordenes_compra.xlsx', sheet_name='Sheet1')",
        "    df_ordenes.head()",
        "Cargue los logs transaccionales desde JSON escribiendo la siguiente sintaxis y ejecútela:",
        "    df_logs = pd.read_json('logs_transacciones.json')",
        "    print('Logs cargados con éxito. Filas cargadas:', len(df_logs))"
    ]
    for step in steps_ex3:
        if step.strip().startswith("df_") or step.strip().startswith("print") or step.strip().startswith("import"):
            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Inches(0.6)
            p_code.paragraph_format.space_after = Pt(4)
            run_c = p_code.add_run(step)
            format_run(run_c, font_name="Consolas", size_pt=9.5, rgb_color=COLOR_PRIMARY_RGB)
        else:
            p_step = doc.add_paragraph(style='List Number')
            p_step.paragraph_format.space_after = Pt(4)
            p_step.paragraph_format.left_indent = Inches(0.4)
            run_s = p_step.add_run(step)
            format_run(run_s, size_pt=10)
            
    add_callout(doc, "✓ CHECKPOINT: El DataFrame df_logs deberá mostrar de forma estructurada las claves del JSON como columnas individuales de la tabla cargada en memoria.")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # 5. EVALUACIÓN Y CRITERIOS
    h5 = doc.add_paragraph()
    r_h5 = h5.add_run("5. EVALUACIÓN DE DESEMPEÑO")
    format_run(r_h5, size_pt=12, bold=True, rgb_color=COLOR_PRIMARY_RGB)
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(6)
    
    p_criteria_desc = doc.add_paragraph()
    p_criteria_desc.paragraph_format.space_after = Pt(6)
    r_criteria_desc = p_criteria_desc.add_run("Para acreditar este laboratorio, el estudiante debe demostrar al instructor:")
    format_run(r_criteria_desc, size_pt=10, italic=True)
    
    criteria = [
        "El Jupyter Notebook en funcionamiento continuo sin errores de importación de librerías.",
        "La correcta carga en memoria y previsualización de los DataFrames df_clientes, df_ordenes y df_logs.",
        "El listado estructural por pantalla de los tres datasets imprimiendo correctamente sus dimensiones y nombres de campos."
    ]
    for cr in criteria:
        p_bullet = doc.add_paragraph(style='List Bullet')
        p_bullet.paragraph_format.space_after = Pt(4)
        p_bullet.paragraph_format.left_indent = Inches(0.4)
        run_b = p_bullet.add_run(cr)
        format_run(run_b, size_pt=10)
        
    # Save document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Guía de Laboratorio guardada exitosamente en {output_path}")

if __name__ == "__main__":
    base_dir = r"C:\Users\ianache\Desktop\DATA\01-DOCUMENTOS\02-PROYECTOS\112-skillvault\PY-101"
    generate_ficha(os.path.join(base_dir, "PY-101-L1-ficha.docx"))
    generate_presentation(os.path.join(base_dir, "PY-101-L1-presentacion.pptx"))
    generate_lab_guide(os.path.join(base_dir, "PY-101-L1-lab.docx"))
    print("PROCESO DE GENERACIÓN COMPLETADO DE MANERA EXITOSA.")
